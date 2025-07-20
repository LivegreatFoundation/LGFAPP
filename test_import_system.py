#!/usr/bin/env python3
"""
Test the blog import system with sample data
Live Great Foundation - Django Blog System
"""

import os
import sys
import django
from openpyxl import Workbook

# Setup Django
sys.path.append('.')
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'blog.settings')
django.setup()

from blogapp.admin import PostResource
from blogapp.models import Category
from taggit.models import Tag

def create_test_import_file():
    """Create a test Excel file with sample blog posts"""
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Import"
    
    # Headers
    headers = ['title', 'content', 'Author', 'category', 'status', 'tags', 'featured', 'trending', 'image']
    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header)
    
    # Test data
    test_posts = [
        {
            'title': 'Test Post: Community Garden Success in Kibera',
            'content': '''<h2>Growing Hope in Urban Slums</h2>
<p>The Live Great Foundation's community garden project in Kibera has transformed a vacant lot into a thriving source of fresh vegetables and community pride.</p>

<h3>Project Highlights</h3>
<ul>
<li>50 families now have access to fresh vegetables</li>
<li>Monthly income increased by KES 3,000 per family</li>
<li>Children learn about nutrition and farming</li>
<li>Community cooperation has strengthened</li>
</ul>

<blockquote>"This garden has changed our lives. My children eat vegetables every day now." - Mary Wanjiku, Community Member</blockquote>

<p>The success of this project demonstrates the power of community-led development initiatives.</p>''',
            'Author': 'Test Author',
            'category': 'Community Development',
            'status': 'published',
            'tags': 'community, agriculture, kibera, urban-farming, nutrition',
            'featured': 'TRUE',
            'trending': 'FALSE',
            'image': ''
        }
    ]
    
    # Add test data
    for row, post in enumerate(test_posts, 2):
        ws.cell(row=row, column=1, value=post['title'])
        ws.cell(row=row, column=2, value=post['content'])
        ws.cell(row=row, column=3, value=post['Author'])
        ws.cell(row=row, column=4, value=post['category'])
        ws.cell(row=row, column=5, value=post['status'])
        ws.cell(row=row, column=6, value=post['tags'])
        ws.cell(row=row, column=7, value=post['featured'])
        ws.cell(row=row, column=8, value=post['trending'])
        ws.cell(row=row, column=9, value=post['image'])
    
    # Save test file
    test_file = 'test_blog_import.xlsx'
    wb.save(test_file)
    return test_file

def test_import_resource():
    """Test the PostResource import functionality"""
    
    print("🧪 Testing Blog Import System")
    print("=" * 50)
    
    # Check if categories exist
    categories = Category.objects.all()
    print(f"📂 Available Categories: {categories.count()}")
    for cat in categories:
        print(f"   • {cat.title}")
    
    # Check if tags exist
    tags = Tag.objects.all()
    print(f"\n🏷️  Available Tags: {tags.count()}")
    print(f"   Sample tags: {', '.join([tag.name for tag in tags[:10]])}")
    
    # Test resource configuration
    resource = PostResource()
    print(f"\n⚙️  Resource Configuration:")
    print(f"   • Model: {resource._meta.model.__name__}")
    print(f"   • Fields: {', '.join(resource._meta.fields)}")
    print(f"   • Import ID Fields: {resource._meta.import_id_fields}")
    
    # Create test file
    print(f"\n📄 Creating test import file...")
    test_file = create_test_import_file()
    print(f"   ✅ Created: {test_file}")
    
    print(f"\n🎯 Next Steps:")
    print(f"   1. Open Django Admin: http://localhost:8000/admin/blogapp/post/")
    print(f"   2. Click 'Download Import Template' to get the full template")
    print(f"   3. Use the Import button to test with {test_file}")
    print(f"   4. Verify the import process works correctly")
    
    return test_file

if __name__ == "__main__":
    test_file = test_import_resource()
    print(f"\n✅ Test setup complete! Test file: {test_file}")
